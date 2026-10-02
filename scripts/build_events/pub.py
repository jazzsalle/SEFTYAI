import openpyxl,json,collections
base="D:/VIBECODING/SafetyAI/seed_data/09. 주관기관 데이터 수신/16. ★상습침수지역 및 피해데이터(재난영향분석과 자료)_260811/수신 데이터/★침수 피해 데이터(2차)/★침수 피해 데이터(2차)/"
def reg(s):
  s=str(s or '')
  return '부산' if s.startswith('부산') else '영천' if '영천시' in s else '인제' if '인제군' in s else None
res={}
wb=openpyxl.load_workbook(base+"(복구 확정) 2011~2025 자연재난 공공시설 피해내역_침수 한정(4,019개).xlsx",read_only=True)
rows=list(wb.active.iter_rows(values_only=True)); h=rows[1]
pub=[dict(zip(h,r)) for r in rows[2:] if reg(r[5])]
for d in pub: d['_reg']=reg(d['피해상세주소'])
res['public']=pub
c=collections.Counter((d['_reg'],d['재난년도'],d['재난명']) for d in pub)
for k,v in sorted(c.items(),key=str): print('PUB',k,v)
for fn,key in (("소상공인_2017~2025년_풍수해보험금지급내역_침수 한정(2,553개).xlsx",'sme'),("주택_2016~2025년_풍수해보험금지급내역_침수 한정(4,433개).xlsx",'house')):
  wb=openpyxl.load_workbook(base+fn,read_only=True)
  rows=list(wb.active.iter_rows(values_only=True)); h=rows[1]
  L=[]
  for r in rows[2:]:
    d=dict(zip(h,r)); g=reg(str(d.get('지역시도') or '')+' '+str(d.get('지역시군구') or '')) or reg(d.get('보상목적물주소'))
    if g:
      # anonymize: drop names & policy no
      for k in ('계약자성명','피보험자성명','증권번호','계약자주소','손해평가인'): d.pop(k,None)
      d['_reg']=g; L.append(d)
  res[key]=L
  c=collections.Counter((d['_reg'],str(d['사고일자'])[:8],d['사고유형']) for d in L)
  for k,v in sorted(c.items()): print(key,k,v)
json.dump(res,open(r"D:/VIBECODING/SafetyAI/scripts/build_events/_work/flood3.json","w",encoding="utf-8"),ensure_ascii=False,default=str)
